#!/usr/bin/env python3
"""Build a story-first PowerPoint without modifying the paper.

The deck reads the paper's machine-readable evidence and existing conceptual
figures, creates slide-native plots, and writes an editable PPTX (text and
layout remain PowerPoint objects; statistical plots are embedded PNG assets).
"""

from __future__ import annotations

import hashlib
import json
import math
import shutil
import statistics
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / ".local" / "pptx"))

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.colors import LinearSegmentedColormap, TwoSlopeNorm
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = ROOT / "presentation"
FIG = OUT / "figures"
PPTX_OUT = OUT / "verified_mode_support_story.pptx"
MANIFEST_OUT = OUT / "deck_manifest.json"

CROSS_JSON = ROOT / "paper" / "figures" / "cross_scale_terminal_endpoint_effects.json"
AUC_JSON = ROOT / "paper" / "figures" / "sustained_auc_effects_qwen05b.json"
TELEMETRY_JSON = ROOT / "paper" / "figures" / "replay_mechanism_telemetry_qwen05b.json"
DIRECT_JSON = ROOT / "paper" / "figures" / "direct_comparator_endpoint_effects.json"
DISCOVERY_JSON = ROOT / "paper" / "figures" / "verified_support_discovery_two_scale_effects.json"

PAPER_COLLAPSE = ROOT / "paper" / "figures" / "modecollapse_story.png"
PAPER_MODEBENCH = ROOT / "paper" / "figures" / "modebench_examples.png"
PAPER_METHOD = ROOT / "paper" / "figures" / "verified_support_story.png"

SLIDE_W = 13.333
SLIDE_H = 7.5

# Deck palette: warm editorial background, dark ink, teal intervention.
BG = "F7F4EE"
INK = "15263D"
MUTED = "617083"
GRID = "DCE2E8"
CARD = "EAF0F5"
WHITE = "FFFFFF"
TEAL = "0D9488"
TEAL_DARK = "08766E"
TEAL_LIGHT = "BFE7E2"
PURPLE = "6D4AE5"
PURPLE_LIGHT = "DDD4FA"
CORAL = "E25B55"
CORAL_LIGHT = "F5D2CE"
GOLD = "D39B37"
BLUE = "3A78C2"
GRAY = "9AA5B1"
FONT = "DejaVu Sans"
MONO = "DejaVu Sans Mono"

DOMAIN_ORDER = ["graph_coloring", "countdown", "python_factors", "mathir", "pantry_plan"]
DOMAIN_LABEL = {
    "graph_coloring": "Graph",
    "countdown": "Countdown",
    "python_factors": "Python",
    "mathir": "MathIR",
    "pantry_plan": "Pantry",
}
MODEL_ORDER = ["Qwen2.5-0.5B", "Falcon3-1B", "Qwen2.5-3B"]
MODEL_COLOR = {
    "Qwen2.5-0.5B": TEAL,
    "Falcon3-1B": PURPLE,
    "Qwen2.5-3B": GOLD,
}


def read_json(path: Path) -> dict:
    with path.open() as f:
        return json.load(f)


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.replace("#", ""))


def mpl_color(hex_color: str) -> str:
    return f"#{hex_color}"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def setup_matplotlib() -> None:
    mpl.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 12,
            "axes.titlesize": 15,
            "axes.labelsize": 12,
            "xtick.labelsize": 10,
            "ytick.labelsize": 11,
            "axes.edgecolor": mpl_color(GRID),
            "axes.labelcolor": mpl_color(INK),
            "xtick.color": mpl_color(MUTED),
            "ytick.color": mpl_color(MUTED),
            "text.color": mpl_color(INK),
            "axes.facecolor": mpl_color(BG),
            "figure.facecolor": mpl_color(BG),
            "savefig.facecolor": mpl_color(BG),
            "grid.color": mpl_color(GRID),
            "grid.linewidth": 0.8,
            "axes.spines.top": False,
            "axes.spines.right": False,
        }
    )


def save_figure(fig: plt.Figure, name: str) -> Path:
    path = FIG / name
    fig.savefig(path, dpi=240, bbox_inches="tight", pad_inches=0.10)
    plt.close(fig)
    return path


def paired_effects(cross: dict) -> list[dict]:
    effects = []
    for row in cross["rows"]:
        for seed, ep in row["per_seed_endpoints"].items():
            d_pass = ep["replay"]["pass8"] - ep["control"]["pass8"]
            d_distinct = ep["replay"]["distinct8"] - ep["control"]["distinct8"]
            effects.append(
                {
                    "model": row["model"],
                    "domain": row["domain"],
                    "seed": int(seed),
                    "d_pass": d_pass,
                    "d_distinct": d_distinct,
                    "d_adjusted": d_distinct - d_pass,
                }
            )
    return effects


def student_t_interval(values: list[float]) -> tuple[float, float]:
    mean = statistics.fmean(values)
    if len(values) < 2:
        return mean, mean
    # Every confirmatory block in this deck has n=5 (df=4).
    tcrit = 2.7764451051977987 if len(values) == 5 else 1.96
    half = tcrit * statistics.stdev(values) / math.sqrt(len(values))
    return mean - half, mean + half


def build_headline_pairs(cross: dict) -> Path:
    effects = paired_effects(cross)
    fig, axes = plt.subplots(1, 3, figsize=(14.6, 4.45), sharex=True, sharey=True)
    jitter = np.linspace(-0.14, 0.14, 5)
    for ax, model in zip(axes, MODEL_ORDER):
        color = mpl_color(MODEL_COLOR[model])
        for yi, domain in enumerate(DOMAIN_ORDER[::-1]):
            vals = [e["d_distinct"] for e in effects if e["model"] == model and e["domain"] == domain]
            ax.scatter(vals, yi + jitter[: len(vals)], s=34, color=color, alpha=0.48, edgecolors="none", zorder=2)
            mean = statistics.fmean(vals)
            lo, hi = student_t_interval(vals)
            ax.plot([lo, hi], [yi, yi], color=mpl_color(INK), lw=2.0, zorder=3)
            ax.scatter([mean], [yi], s=82, marker="D", color=color, edgecolors=mpl_color(INK), linewidths=0.8, zorder=4)
        ax.axvline(0, color=mpl_color(CORAL), lw=1.5, ls="--")
        ax.grid(axis="x")
        ax.set_title(model, fontweight="bold", pad=10)
        ax.text(0.97, 0.05, "25 / 25  > 0", transform=ax.transAxes, ha="right", va="bottom", fontsize=11, fontweight="bold", color=color)
        ax.set_xlim(-0.08, 2.55)
        ax.set_xlabel("Replay − control  Δdistinct@8")
    axes[0].set_yticks(range(5), [DOMAIN_LABEL[d] for d in DOMAIN_ORDER[::-1]])
    fig.subplots_adjust(left=0.09, right=0.99, top=0.91, bottom=0.18, wspace=0.12)
    return save_figure(fig, "headline_75_pairs.png")


def build_falcon_dumbbells(cross: dict) -> Path:
    rows = {r["domain"]: r for r in cross["rows"] if r["model"] == "Falcon3-1B"}
    fig, axes = plt.subplots(1, 5, figsize=(14.7, 4.0), sharey=True)
    for ax, domain in zip(axes, DOMAIN_ORDER):
        row = rows[domain]
        controls, replays = [], []
        for ep in row["per_seed_endpoints"].values():
            c = ep["control"]["distinct8"]
            r = ep["replay"]["distinct8"]
            controls.append(c)
            replays.append(r)
            ax.plot([0, 1], [c, r], color=mpl_color(TEAL_LIGHT), lw=2.0, zorder=1)
            ax.scatter([0], [c], s=33, color=mpl_color(GRAY), zorder=2)
            ax.scatter([1], [r], s=38, color=mpl_color(TEAL), zorder=2)
        cm, rm = statistics.fmean(controls), statistics.fmean(replays)
        ax.plot([-0.12, 0.12], [cm, cm], color=mpl_color(INK), lw=3)
        ax.plot([0.88, 1.12], [rm, rm], color=mpl_color(TEAL_DARK), lw=3)
        ax.set_xlim(-0.28, 1.28)
        ax.set_xticks([0, 1], ["Control", "Replay"])
        ax.set_title(DOMAIN_LABEL[domain], fontweight="bold")
        ax.grid(axis="y")
        ax.text(0.5, 0.96, f"+{rm-cm:.3f}", transform=ax.transAxes, ha="center", va="top", color=mpl_color(TEAL_DARK), fontsize=12, fontweight="bold")
    axes[0].set_ylabel("Terminal distinct verified modes in 8 samples")
    axes[0].set_ylim(-0.05, 3.55)
    fig.subplots_adjust(left=0.07, right=0.995, top=0.91, bottom=0.17, wspace=0.18)
    return save_figure(fig, "falcon_paired_endpoints.png")


def build_accuracy_breadth_scatter(cross: dict) -> Path:
    effects = paired_effects(cross)
    fig, ax = plt.subplots(figsize=(10.8, 5.45))
    x_max = max(e["d_pass"] for e in effects) * 1.08
    y_max = max(e["d_distinct"] for e in effects) * 1.08
    line_max = max(x_max, y_max)
    ax.fill_between([0, line_max], [0, line_max], [line_max, line_max], color=mpl_color(TEAL_LIGHT), alpha=0.45, zorder=0)
    ax.plot([0, line_max], [0, line_max], color=mpl_color(MUTED), lw=1.6, ls="--")
    markers = {"graph_coloring": "o", "countdown": "s", "python_factors": "^", "mathir": "D", "pantry_plan": "P"}
    for model in MODEL_ORDER:
        color = mpl_color(MODEL_COLOR[model])
        for domain in DOMAIN_ORDER:
            pts = [e for e in effects if e["model"] == model and e["domain"] == domain]
            ax.scatter(
                [p["d_pass"] for p in pts],
                [p["d_distinct"] for p in pts],
                s=60,
                marker=markers[domain],
                color=color,
                alpha=0.72,
                edgecolors=mpl_color(WHITE),
                linewidths=0.7,
            )
    ax.axvline(0, color=mpl_color(GRID), lw=1)
    ax.axhline(0, color=mpl_color(GRID), lw=1)
    ax.set_xlim(-0.02, x_max)
    ax.set_ylim(-0.05, y_max)
    ax.set_xlabel("Δpass@8  (added probability of ≥1 correct sample)", fontweight="bold")
    ax.set_ylabel("Δdistinct@8  (added verified outcomes)", fontweight="bold")
    ax.grid(alpha=0.75)
    ax.text(0.97, 0.92, "breadth beyond\nadded correctness", transform=ax.transAxes, ha="right", va="top", fontsize=12, color=mpl_color(TEAL_DARK), fontweight="bold")
    ax.text(0.97, 0.08, "chiefly an\naccuracy rescue", transform=ax.transAxes, ha="right", va="bottom", fontsize=11, color=mpl_color(MUTED))
    model_handles = [mpl.lines.Line2D([0], [0], marker="o", color="none", markerfacecolor=mpl_color(MODEL_COLOR[m]), markersize=9, label=m) for m in MODEL_ORDER]
    domain_handles = [mpl.lines.Line2D([0], [0], marker=markers[d], color=mpl_color(INK), lw=0, markersize=7, label=DOMAIN_LABEL[d]) for d in DOMAIN_ORDER]
    leg1 = ax.legend(handles=model_handles, loc="upper left", frameon=False, ncol=1, fontsize=10)
    ax.add_artist(leg1)
    ax.legend(handles=domain_handles, loc="lower left", frameon=False, ncol=3, fontsize=9, columnspacing=1.0, handletextpad=0.4)
    fig.tight_layout()
    return save_figure(fig, "accuracy_vs_breadth_effects.png")


def discovery_model_summary(discovery: dict, model: str) -> dict:
    cells = [c for c in discovery["cells"] if c["model"] == model]
    pass_values = []
    breadth_values = []
    for cell in cells:
        for effects in cell["per_seed_effects"].values():
            pass_values.append(effects["terminal_sampled_pass8"])
            breadth_values.append(effects["terminal_sampled_excess8"])
    return {
        "pairs": len(breadth_values),
        "mean_pass": statistics.fmean(pass_values),
        "mean_breadth": statistics.fmean(breadth_values),
        "positive_domains": sum(
            cell["summaries"]["terminal_sampled_excess8"]["mean"] > 0 for cell in cells
        ),
    }


def build_discovery_effects(discovery: dict) -> Path:
    models = ["Qwen 0.5B", "Falcon 1B"]
    colors = [mpl_color(TEAL), mpl_color(PURPLE)]
    cells = {(c["model"], c["domain"]): c for c in discovery["cells"]}
    domains_rev = DOMAIN_ORDER[::-1]
    fig, axes = plt.subplots(1, 2, figsize=(12.8, 4.7), sharex=True, sharey=True)
    jitter = np.linspace(-0.12, 0.12, 5)
    for ax, model, color in zip(axes, models, colors):
        for yi, domain in enumerate(domains_rev):
            cell = cells[(model, domain)]
            values = list(
                cell["summaries"]["terminal_sampled_excess8"]["per_seed"].values()
            )
            summary = cell["summaries"]["terminal_sampled_excess8"]
            lo, hi = summary["student_t_95"]
            ax.scatter(
                values,
                yi + jitter[: len(values)],
                s=38,
                facecolors=mpl_color(BG),
                edgecolors=color,
                linewidths=1.4,
                zorder=2,
            )
            ax.plot([lo, hi], [yi, yi], color=mpl_color(INK), lw=1.8, zorder=3)
            ax.scatter(
                [summary["mean"]],
                [yi],
                s=74,
                marker="D",
                color=color,
                edgecolors=mpl_color(INK),
                linewidths=0.7,
                zorder=4,
            )
            ax.text(1.47, yi, f"n={cell['n']}", ha="right", va="center", fontsize=9, color=mpl_color(MUTED))
        ax.axvline(0, color=mpl_color(CORAL), lw=1.4, ls="--")
        ax.grid(axis="x")
        ax.set_xlim(-0.27, 1.52)
        ax.set_title(model, fontweight="bold", pad=10)
        ax.set_xlabel("Δ correctness-adjusted breadth@8")
    axes[0].set_yticks(np.arange(5), [DOMAIN_LABEL[d] for d in domains_rev])
    fig.subplots_adjust(left=0.12, right=0.99, top=0.90, bottom=0.18, wspace=0.12)
    return save_figure(fig, "verified_support_discovery_effects.png")


def build_auc_effects(auc: dict) -> Path:
    metric_specs = [
        ("normalized_auc_pass8", "Correctness AUC", BLUE),
        ("normalized_auc_distinct8", "Raw breadth AUC", TEAL),
        ("normalized_auc_adjusted_breadth8", "Adjusted breadth AUC", PURPLE),
    ]
    cells = {c["domain"]: c for c in auc["cells"]}
    fig, axes = plt.subplots(1, 3, figsize=(14.6, 4.4), sharey=True)
    domains_rev = DOMAIN_ORDER[::-1]
    for ax, (metric, title, color_hex) in zip(axes, metric_specs):
        color = mpl_color(color_hex)
        means, los, his = [], [], []
        for domain in domains_rev:
            summary = cells[domain]["effects"][metric]
            means.append(summary["mean"])
            los.append(summary["student_t_95"][0])
            his.append(summary["student_t_95"][1])
        y = np.arange(len(domains_rev))
        ax.axvline(0, color=mpl_color(CORAL), lw=1.3, ls="--")
        ax.errorbar(means, y, xerr=[np.array(means) - np.array(los), np.array(his) - np.array(means)], fmt="D", ms=7, color=color, ecolor=mpl_color(INK), elinewidth=1.8, capsize=3)
        ax.set_title(title, fontweight="bold")
        ax.set_xlabel("Replay − control")
        ax.grid(axis="x")
        xlo = min(min(los), 0)
        xhi = max(his)
        pad = max(0.05, (xhi - xlo) * 0.15)
        ax.set_xlim(xlo - pad, xhi + pad)
    axes[0].set_yticks(np.arange(5), [DOMAIN_LABEL[d] for d in domains_rev])
    fig.subplots_adjust(left=0.09, right=0.99, top=0.91, bottom=0.18, wspace=0.16)
    return save_figure(fig, "sustained_auc_effects.png")


def build_telemetry(telemetry: dict) -> Path:
    domains = [telemetry["domains"][d] for d in DOMAIN_ORDER]
    labels = [d["label"] for d in domains]
    occ = [d["late_occupancy"]["mean"] for d in domains]
    times = [d["paired_optimizer_time_percent_effect"]["mean"] for d in domains]
    fig = plt.figure(figsize=(14.5, 4.3))
    gs = fig.add_gridspec(1, 3, width_ratios=[1.15, 0.72, 1.15], wspace=0.34)
    ax1 = fig.add_subplot(gs[0, 0])
    ax2 = fig.add_subplot(gs[0, 1])
    ax3 = fig.add_subplot(gs[0, 2])
    x = np.arange(5)
    ax1.bar(x, occ, color=mpl_color(TEAL), alpha=0.88, width=0.7)
    for i, d in enumerate(domains):
        vals = list(d["late_occupancy"]["seed_values"].values())
        ax1.scatter(np.full(len(vals), i) + np.linspace(-0.12, 0.12, len(vals)), vals, s=22, color=mpl_color(INK), alpha=0.55)
    ax1.axhline(telemetry["capacity"], color=mpl_color(CORAL), ls="--", lw=1.5)
    ax1.text(4.45, telemetry["capacity"] - 0.4, "cap = 16", ha="right", va="top", color=mpl_color(CORAL), fontsize=10, fontweight="bold")
    ax1.set_xticks(x, labels, rotation=20, ha="right")
    ax1.set_ylabel("Late bank occupancy (modes)")
    ax1.set_title("Small prompt-local banks", fontweight="bold")
    ax1.set_ylim(0, 17.5)
    ax1.grid(axis="y")

    hit = telemetry["capacity_hit_fraction"]
    ax2.pie([hit, 1 - hit], colors=[mpl_color(CORAL), mpl_color(CARD)], startangle=90, wedgeprops={"width": 0.24, "edgecolor": mpl_color(BG)})
    ax2.text(0, 0.08, f"{100*hit:.1f}%", ha="center", va="center", fontsize=25, fontweight="bold", color=mpl_color(INK))
    ax2.text(0, -0.22, "updates hit cap", ha="center", va="center", fontsize=10, color=mpl_color(MUTED))
    ax2.set_title("Capacity rarely binds", fontweight="bold", pad=12)
    ax2.text(0.5, -0.10, f"{telemetry['capacity_hit_updates']:,} / {telemetry['replay_updates']:,}\nreplay updates", transform=ax2.transAxes, ha="center", va="top", fontsize=10, color=mpl_color(MUTED))

    colors = [mpl_color(TEAL) if v < 5 else mpl_color(PURPLE) for v in times]
    ax3.bar(x, times, color=colors, alpha=0.88, width=0.7)
    for i, d in enumerate(domains):
        vals = list(d["paired_optimizer_time_percent_effect"]["seed_values"].values())
        ax3.scatter(np.full(len(vals), i) + np.linspace(-0.12, 0.12, len(vals)), vals, s=22, color=mpl_color(INK), alpha=0.5)
    ax3.axhline(0, color=mpl_color(MUTED), lw=1)
    ax3.set_xticks(x, labels, rotation=20, ha="right")
    ax3.set_ylabel("Optimizer-time effect (%)")
    ax3.set_title("Descriptive compute overhead", fontweight="bold")
    ax3.grid(axis="y")
    fig.subplots_adjust(left=0.06, right=0.99, top=0.89, bottom=0.22)
    return save_figure(fig, "replay_mechanism_telemetry.png")


def build_comparator_heatmap(cross: dict, direct: dict) -> Path:
    models = ["Qwen2.5-0.5B", "Falcon3-1B"]
    methods = ["ReplayDr.GRPO", "RLEP-Dr", "UCPO", "GRPO"]
    cols = [(m, d) for m in models for d in DOMAIN_ORDER]
    values = np.full((len(methods), len(cols)), np.nan)
    intervals: dict[tuple[int, int], tuple[float, float]] = {}
    replay_rows = {(r["model"], r["domain"]): r for r in cross["rows"]}
    direct_cells = {(c["model"], c["domain"]): c for c in direct["cells"]}
    direct_key = {"RLEP-Dr": "rlep_dr", "UCPO": "ucpo", "GRPO": "grpo"}
    for j, key in enumerate(cols):
        r = replay_rows[key]
        values[0, j] = r["summaries"]["adjusted_breadth8"]["mean"]
        intervals[(0, j)] = tuple(r["summaries"]["adjusted_breadth8"]["student_t_95"])
        for i, method in enumerate(methods[1:], start=1):
            rec = direct_cells[key]["methods"].get(direct_key[method])
            if rec and rec["n"] == 5:
                values[i, j] = rec["summaries"]["adjusted_breadth8"]["mean"]
                intervals[(i, j)] = tuple(rec["summaries"]["adjusted_breadth8"]["student_t_95"])
    cmap = LinearSegmentedColormap.from_list("deck_div", [mpl_color(CORAL), mpl_color(BG), mpl_color(TEAL_DARK)])
    norm = TwoSlopeNorm(vmin=-0.15, vcenter=0, vmax=1.55)
    fig, ax = plt.subplots(figsize=(14.5, 4.25))
    masked = np.ma.masked_invalid(values)
    im = ax.imshow(masked, cmap=cmap, norm=norm, aspect="auto")
    ax.set_yticks(np.arange(len(methods)), methods)
    ax.set_xticks(np.arange(len(cols)), [DOMAIN_LABEL[d] for _, d in cols], rotation=24, ha="right")
    ax.set_xticks(np.arange(-0.5, len(cols), 1), minor=True)
    ax.set_yticks(np.arange(-0.5, len(methods), 1), minor=True)
    ax.grid(which="minor", color=mpl_color(WHITE), linewidth=2)
    ax.tick_params(which="minor", bottom=False, left=False)
    ax.axvline(4.5, color=mpl_color(INK), lw=2.2)
    ax.text(2, -0.82, "Qwen2.5-0.5B", ha="center", va="bottom", fontsize=12, fontweight="bold")
    ax.text(7, -0.82, "Falcon3-1B", ha="center", va="bottom", fontsize=12, fontweight="bold")
    for i in range(values.shape[0]):
        for j in range(values.shape[1]):
            if np.isnan(values[i, j]):
                ax.text(j, i, "n<5", ha="center", va="center", fontsize=9, color=mpl_color(MUTED))
                continue
            lo, hi = intervals[(i, j)]
            sig = " •" if lo > 0 or hi < 0 else ""
            color = WHITE if abs(values[i, j]) > 0.38 else INK
            ax.text(j, i, f"{values[i,j]:+.2f}{sig}", ha="center", va="center", fontsize=9.5, fontweight="bold" if i == 0 else "normal", color=mpl_color(color))
    cbar = fig.colorbar(im, ax=ax, orientation="horizontal", fraction=0.08, pad=0.22, aspect=45)
    cbar.set_label("Mean Δ(correctness-adjusted breadth@8) vs matched Dr.GRPO", fontsize=10)
    fig.subplots_adjust(left=0.14, right=0.99, top=0.87, bottom=0.31)
    return save_figure(fig, "direct_comparator_heatmap.png")


def copy_conceptual_assets() -> dict[str, Path]:
    assets = {
        "collapse": (PAPER_COLLAPSE, FIG / "paper_modecollapse_story.png"),
        "modebench": (PAPER_MODEBENCH, FIG / "paper_modebench_examples.png"),
        "method": (PAPER_METHOD, FIG / "paper_verified_replay_mechanism.png"),
    }
    out = {}
    for key, (src, dst) in assets.items():
        shutil.copy2(src, dst)
        out[key] = dst
    return out


def add_bg(slide) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_W), Inches(SLIDE_H))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(BG)
    shape.line.fill.background()
    slide.shapes._spTree.remove(shape._element)
    slide.shapes._spTree.insert(2, shape._element)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    color: str = INK,
    bold: bool = False,
    font: str = FONT,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    margin: float = 0.03,
    line_spacing: float = 1.0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.line_spacing = line_spacing
    for run in p.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return box


def add_rich_text(slide, runs: list[dict], x: float, y: float, w: float, h: float, *, size: float = 18, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    for spec in runs:
        run = p.add_run()
        run.text = spec["text"]
        run.font.name = spec.get("font", FONT)
        run.font.size = Pt(spec.get("size", size))
        run.font.bold = spec.get("bold", False)
        run.font.color.rgb = rgb(spec.get("color", INK))
    return box


def add_card(slide, x: float, y: float, w: float, h: float, *, fill: str = CARD, line: str | None = None, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def add_title(slide, title: str, number: int, kicker: str = "MODE COLLAPSE UNDER GRPO", subtitle: str | None = None) -> None:
    add_text(slide, kicker, 0.48, 0.24, 6.3, 0.24, size=9.5, color=TEAL_DARK, bold=True)
    add_text(slide, f"{number:02d}", 12.28, 0.22, 0.55, 0.28, size=10, color=MUTED, bold=True, align=PP_ALIGN.RIGHT)
    add_text(slide, title, 0.46, 0.52, 12.25, 0.56, size=25.5, color=INK, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.48, 1.08, 12.0, 0.34, size=12.5, color=MUTED)


def add_source(slide, source: str) -> None:
    add_text(slide, source, 0.48, 7.22, 12.15, 0.18, size=7.6, color=MUTED)


def add_picture_contain(slide, path: Path, x: float, y: float, w: float, h: float):
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    return slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))


def add_bullet_list(slide, items: list[tuple[str, str]], x: float, y: float, w: float, h: float, *, size=17, bullet_color=TEAL):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for idx, (head, body) in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = Pt(11)
        p.level = 0
        r = p.add_run()
        r.text = "●  "
        r.font.name = FONT
        r.font.size = Pt(size - 2)
        r.font.color.rgb = rgb(bullet_color)
        r = p.add_run()
        r.text = head
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = rgb(INK)
        r = p.add_run()
        r.text = body
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.color.rgb = rgb(MUTED)
    return box


def add_notes(slide, text: str) -> None:
    try:
        slide.notes_slide.notes_text_frame.text = text
    except Exception:
        pass


def add_arrow(slide, x1: float, y1: float, x2: float, y2: float, color: str = MUTED, width: float = 2.0, dashed: bool = False):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    line.line.end_arrowhead = True
    if dashed:
        line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return line


def add_mode_circle(slide, x: float, y: float, d: float, label: str, fill: str, alpha_hint: bool = False):
    circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    circ.fill.solid()
    circ.fill.fore_color.rgb = rgb(fill)
    circ.line.color.rgb = rgb(WHITE if not alpha_hint else GRID)
    circ.line.width = Pt(1.5)
    add_text(slide, label, x, y + 0.01, d, d - 0.02, size=max(12, 9 * d), color=WHITE if not alpha_hint else MUTED, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)


def build_pptx(assets: dict[str, Path], plots: dict[str, Path], cross: dict, discovery: dict) -> tuple[Presentation, list[dict]]:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = "When Eight Samples Become Eight Copies"
    prs.core_properties.subject = "A story-first presentation of verified mode support in RLVR"
    prs.core_properties.author = "Anonymous Authors"
    blank = prs.slide_layouts[6]
    manifest_slides: list[dict] = []

    def new_slide(title: str, source: str, notes: str = ""):
        slide = prs.slides.add_slide(blank)
        add_bg(slide)
        manifest_slides.append({"number": len(prs.slides), "title": title, "source": source})
        add_notes(slide, notes or f"Source: {source}")
        return slide

    # 1 — title
    slide = new_slide(
        "When eight samples become eight copies",
        "Paper abstract and Figure 1",
        "Opening: Binary-reward RL can solve the task while silently erasing the set of solutions we may need at inference time.",
    )
    add_text(slide, "VERIFIED MODE SUPPORT IN RLVR", 0.58, 0.55, 5.8, 0.3, size=11, color=TEAL_DARK, bold=True)
    add_text(slide, "When eight samples\nbecome eight copies", 0.58, 1.18, 7.7, 1.62, size=38, color=INK, bold=True, line_spacing=0.88)
    add_text(slide, "Measuring—and preserving—the support hidden by binary correctness", 0.62, 3.04, 7.3, 0.68, size=19, color=MUTED)
    add_rich_text(
        slide,
        [
            {"text": "THESIS  ", "bold": True, "color": TEAL_DARK, "size": 11},
            {"text": "Correctness can rise while useful proposal support collapses.", "bold": True, "size": 17},
        ],
        0.62,
        4.23,
        7.3,
        0.52,
    )
    add_text(slide, "Measure  →  Retain  →  Discover", 0.62, 5.0, 6.8, 0.46, size=21, color=PURPLE, bold=True)
    for i in range(8):
        add_mode_circle(slide, 9.0 + (i % 4) * 0.72, 1.45 + (i // 4) * 0.78, 0.48, str(i + 1), TEAL if i < 4 else PURPLE)
    for i in range(8):
        add_arrow(slide, 9.25 + (i % 4) * 0.72, 1.68 + (i // 4) * 0.78, 10.55, 4.15, color=GRID, width=1.1)
    add_mode_circle(slide, 9.76, 3.75, 1.55, "1 mode", PURPLE)
    add_text(slide, "8 nominal attempts", 8.83, 2.98, 3.35, 0.3, size=12, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Anonymous Authors", 0.62, 6.72, 4.5, 0.3, size=11, color=MUTED)

    # 2 — motivating failure
    slide = new_slide(
        "Accuracy can look perfect while support disappears",
        "paper/figures/modecollapse_story.png",
        "Say: Both terminal policies are 32/32 correct. Binary evaluation calls them equal; executed mode identity reveals that one collapsed to a single solution.",
    )
    add_title(slide, "Accuracy can look perfect while support disappears", 2, subtitle="A mechanically selected Qwen2.5-3B Graph Coloring pair: both arms remain 32 / 32 correct.")
    add_picture_contain(slide, assets["collapse"], 0.52, 1.43, 12.25, 4.55)
    add_card(slide, 0.68, 6.18, 12.0, 0.73, fill=INK)
    add_rich_text(slide, [{"text": "Matched Dr.GRPO: ", "bold": True, "color": WHITE}, {"text": "4 → 1 observed modes", "bold": True, "color": CORAL_LIGHT}, {"text": "     ReplayDr.GRPO: ", "bold": True, "color": WHITE}, {"text": "4 → 5 observed modes", "bold": True, "color": TEAL_LIGHT}], 0.92, 6.36, 11.5, 0.3, size=17, align=PP_ALIGN.CENTER)
    add_source(slide, "Source: paper Figure 1 · fixed-seed illustration; registered paired estimates follow")

    # 3 — measurement
    slide = new_slide(
        "Binary correctness hides answer identity",
        "paper/figures/modebench_examples.png",
        "Transition: To train against collapse, the evaluator must identify which successful execution occurred—not merely whether the answer passed.",
    )
    add_title(slide, "Binary correctness hides answer identity", 3, subtitle="ModeBench uses the accepting execution to return a canonical outcome key.")
    add_picture_contain(slide, assets["modebench"], 0.48, 1.43, 9.85, 5.47)
    add_card(slide, 10.52, 1.55, 2.32, 4.95, fill=INK)
    add_text(slide, "Two metrics,\ntwo questions", 10.76, 1.86, 1.85, 0.7, size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "pass@8", 10.77, 2.86, 1.82, 0.36, size=20, color=TEAL_LIGHT, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Did any sample\nsucceed?", 10.77, 3.28, 1.82, 0.64, size=13, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, "distinct@8", 10.67, 4.23, 2.02, 0.36, size=20, color=PURPLE_LIGHT, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "How many different\nverified outcomes?", 10.69, 4.65, 1.98, 0.72, size=13, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, "Token entropy cannot answer the second question.", 10.75, 5.67, 1.85, 0.52, size=11.5, color=CORAL_LIGHT, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "Source: paper Figure 2 · execution-defined keys; no text clustering or gold mode catalogue")

    # 4 — why collapse
    slide = new_slide(
        "Why GRPO concentrates correct modes",
        "Paper Section 2 and Appendix A",
        "Say: Equal reward does not produce equal pressure. Frequent correct modes appear in more sampled updates; absent modes receive no gradient. Replay changes that asymmetry after discovery.",
    )
    add_title(slide, "Why GRPO concentrates correct modes", 4, subtitle="Equal binary reward does not imply equal learning pressure across correct outcomes.")
    add_card(slide, 0.54, 1.52, 5.95, 4.85, fill=CORAL_LIGHT)
    add_text(slide, "ON-POLICY BINARY REWARD", 0.82, 1.83, 5.4, 0.28, size=10.5, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Sampled winners get reinforced", 0.82, 2.17, 5.4, 0.4, size=21, color=INK, bold=True, align=PP_ALIGN.CENTER)
    for i, (lab, fill) in enumerate([("A", PURPLE), ("B", TEAL), ("C", GOLD)]):
        add_mode_circle(slide, 1.10 + i * 1.55, 3.06, 0.72, lab, fill)
    add_arrow(slide, 1.46, 3.86, 2.58, 4.72, color=CORAL, width=2.0)
    add_arrow(slide, 3.01, 3.86, 2.88, 4.72, color=GRID, width=1.4, dashed=True)
    add_arrow(slide, 4.56, 3.86, 3.18, 4.72, color=GRID, width=1.4, dashed=True)
    add_mode_circle(slide, 2.18, 4.63, 1.4, "A", PURPLE)
    add_text(slide, "A correct mode absent from the sampled group receives no update.", 0.95, 5.34, 5.1, 0.78, size=13.5, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
    add_card(slide, 6.83, 1.52, 5.95, 4.85, fill=TEAL_LIGHT)
    add_text(slide, "RECURRENT VERIFIED REPLAY", 7.12, 1.83, 5.4, 0.28, size=10.5, color=TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Every banked mode keeps a gradient", 7.10, 2.17, 5.4, 0.4, size=21, color=INK, bold=True, align=PP_ALIGN.CENTER)
    for i, (lab, fill) in enumerate([("A", PURPLE), ("B", TEAL), ("C", GOLD)]):
        add_mode_circle(slide, 7.38 + i * 1.55, 3.06, 0.72, lab, fill)
        add_arrow(slide, 7.74 + i * 1.55, 3.86, 8.36 + i * 0.82, 4.72, color=TEAL_DARK, width=2.0)
        add_mode_circle(slide, 7.95 + i * 0.82, 4.66, 0.78, lab, fill)
    add_text(slide, "Post-discovery guarantee: recurrent positive replay prevents banked-mode extinction.", 7.18, 5.28, 5.25, 0.86, size=13.5, color=TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "Theory boundary: fixed-bank categorical assumptions; protects discovered, verified, retained modes only")

    # 5 — unified method
    slide = new_slide(
        "One execution key repairs two different failures",
        "paper/figures/verified_support_story.png",
        "Say: Retention and discovery are distinct bottlenecks. Replay revisits known keys uniformly; verified-support Semantic-MaxEnt creates bounded pressure toward rare validator-confirmed support.",
    )
    add_title(slide, "One execution key repairs two different failures", 5, subtitle="Retain what the policy has found; add verified pressure toward what it still misses.")
    add_picture_contain(slide, assets["method"], 0.58, 1.46, 12.18, 4.84)
    add_card(slide, 0.74, 6.38, 12.0, 0.48, fill=INK)
    add_rich_text(slide, [{"text": "RETENTION  ", "bold": True, "color": TEAL_LIGHT}, {"text": "one exemplar per key + uniform replay", "color": WHITE}, {"text": "     DISCOVERY  ", "bold": True, "color": PURPLE_LIGHT}, {"text": "verified proposals + bounded rarity credit", "color": WHITE}], 0.98, 6.49, 11.5, 0.26, size=13.5, align=PP_ALIGN.CENTER)
    add_source(slide, "Both paths use validator-confirmed keys · no unseen-mode oracle or gold support catalogue")

    # 6 — experimental logic
    slide = new_slide(
        "Two questions, two matched tests",
        "Paper experimental design and evidence ledger",
        "Say: The retention test is complete and confirmatory. The discovery test is a smaller exploratory bundled contrast. Keeping those evidential roles separate is central to the clean story.",
    )
    add_title(slide, "Two questions, two matched tests", 6, subtitle="The deck keeps the confirmatory retention result separate from the exploratory discovery extension.")
    add_card(slide, 0.58, 1.53, 12.18, 2.18, fill=TEAL_LIGHT)
    add_text(slide, "RETENTION · CONFIRMATORY", 0.88, 1.81, 3.0, 0.26, size=10.5, color=TEAL_DARK, bold=True)
    add_text(slide, "Does uniform mode replay preserve what was found?", 0.88, 2.18, 5.3, 0.44, size=20, color=INK, bold=True)
    add_text(slide, "ReplayDr.GRPO  −  matched Dr.GRPO", 0.90, 2.79, 4.9, 0.32, size=14, color=TEAL_DARK, bold=True)
    add_text(slide, "3 models  ×  5 domains  ×  5 paired seeds", 6.18, 1.94, 4.25, 0.38, size=17, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "75", 10.63, 1.72, 1.55, 0.78, size=42, color=TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "terminal pairs", 10.55, 2.50, 1.72, 0.32, size=12, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Complete at n=5 in every model–domain cell", 6.25, 2.72, 4.15, 0.32, size=12.5, color=MUTED, align=PP_ALIGN.CENTER)

    add_card(slide, 0.58, 3.96, 12.18, 2.20, fill=PURPLE_LIGHT)
    add_text(slide, "DISCOVERY · EXPLORATORY BUNDLE", 0.88, 4.24, 3.2, 0.26, size=10.5, color=PURPLE, bold=True)
    add_text(slide, "Can verified discovery add support beyond replay?", 0.88, 4.61, 5.3, 0.44, size=20, color=INK, bold=True)
    add_text(slide, "Semantic-MaxEnt + replay  −  replay", 0.90, 5.22, 4.9, 0.32, size=14, color=PURPLE, bold=True)
    add_text(slide, "2 scales  ×  5 domains", 6.18, 4.34, 4.25, 0.38, size=17, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "49", 10.63, 4.12, 1.55, 0.78, size=42, color=PURPLE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "matched endpoints", 10.38, 4.90, 2.05, 0.32, size=12, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "15 / 15 invariant-valid · 11 / 15 complete the full mechanism chain", 6.22, 5.14, 4.2, 0.54, size=12.3, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_card(slide, 1.04, 6.48, 11.26, 0.46, fill=WHITE, line=GRID)
    add_text(slide, "Same prompts and paired initialization · report pass@8, distinct@8, and distinct@8 − pass@8 · no model/domain pooling", 1.25, 6.57, 10.84, 0.25, size=11.5, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "Evidence roles are explicit: 75-pair retention core; 49-endpoint discovery bundle; Qwen3B discovery not analyzed")

    # 7 — retention headline
    slide = new_slide(
        "Retention is the strongest result: 75 / 75",
        str(CROSS_JSON.relative_to(ROOT)),
        "Say: This is the clean headline. Every paired seed at every scale and domain has more raw verified outcomes under replay than under the matched control. Do not pool across domains.",
    )
    add_title(slide, "Retention is the strongest result: 75 / 75", 7, subtitle="Replay raises terminal distinct@8 for every paired seed at every model scale.")
    add_picture_contain(slide, plots["headline"], 0.52, 1.40, 12.30, 5.52)
    add_source(slide, "Source: cross_scale_terminal_endpoint_effects.json · no domain pooling; paired n=5 within each cell")

    # 8 — accuracy adjustment
    slide = new_slide(
        "The honest split: extra support vs. accuracy rescue",
        str(CROSS_JSON.relative_to(ROOT)),
        "Say: Raw breadth is uniformly higher, but not every gain is independent of correctness. Distance above the diagonal is the extra-support effect. This distinction makes the claim defensible.",
    )
    add_title(slide, "The honest split: extra support vs. accuracy rescue", 8, subtitle="Distance above the diagonal is Δdistinct@8 − Δpass@8: verified modes beyond the first success.")
    add_picture_contain(slide, plots["scatter"], 0.56, 1.42, 8.84, 5.55)
    add_card(slide, 9.64, 1.72, 3.10, 4.74, fill=INK)
    add_text(slide, "RAW BREADTH", 9.94, 2.05, 2.52, 0.24, size=10.5, color=TEAL_LIGHT, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "75 / 75", 9.94, 2.40, 2.52, 0.62, size=31, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "paired runs improve", 9.94, 3.00, 2.52, 0.32, size=13, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, "BEYOND ACCURACY", 9.94, 3.73, 2.52, 0.24, size=10.5, color=PURPLE_LIGHT, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Graph · Countdown · Pantry", 9.91, 4.10, 2.58, 0.58, size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "positive domain intervals\nat both Qwen scales", 9.94, 4.68, 2.52, 0.60, size=12.5, color=TEAL_LIGHT, align=PP_ALIGN.CENTER)
    add_text(slide, "Python + MathIR are chiefly accuracy rescues; adjusted effects remain unresolved.", 9.92, 5.48, 2.58, 0.66, size=11.5, color=CORAL_LIGHT, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "Interpretation: replay reliably retains raw modes; independent excess breadth is domain-dependent")

    # 9 — discovery extension
    qwen_summary = discovery_model_summary(discovery, "Qwen 0.5B")
    falcon_summary = discovery_model_summary(discovery, "Falcon 1B")
    slide = new_slide(
        "Discovery adds breadth beyond replay—unevenly",
        str(DISCOVERY_JSON.relative_to(ROOT)),
        "Say: This is a bundled exploratory extension, not the 75-pair core. Mean adjusted breadth is positive at both analyzed scales, driven most clearly by Pantry and Falcon Python; most domain intervals cross zero.",
    )
    add_title(slide, "Discovery adds breadth beyond replay—unevenly", 9, subtitle="Verified-support Semantic-MaxEnt + replay versus replay alone; correctness-adjusted breadth after eight passes.")
    add_picture_contain(slide, plots["discovery"], 0.50, 1.48, 8.86, 5.34)
    add_card(slide, 9.58, 1.66, 3.16, 4.92, fill=PURPLE_LIGHT)
    add_text(slide, "EXPLORATORY BUNDLE", 9.88, 1.99, 2.56, 0.25, size=10.5, color=PURPLE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "49", 9.88, 2.31, 2.56, 0.62, size=34, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "matched endpoints", 9.88, 2.91, 2.56, 0.28, size=12.5, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Mean Δ adjusted breadth", 9.84, 3.48, 2.65, 0.26, size=11.5, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, f"Qwen 0.5B   {qwen_summary['mean_breadth']:+.3f}\nFalcon 1B     {falcon_summary['mean_breadth']:+.3f}", 9.88, 3.82, 2.56, 0.74, size=17, color=PURPLE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "11 / 15", 9.88, 4.81, 2.56, 0.42, size=23, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "cells complete proposal →\npressure → replay", 9.88, 5.22, 2.56, 0.58, size=12.5, color=INK, align=PP_ALIGN.CENTER)
    add_text(slide, "Bundled contrast · Qwen3B not analyzed · most domain intervals cross zero", 9.84, 5.96, 2.66, 0.44, size=10.8, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "Source: 49 integrity-valid smaller-model pairs; Falcon Countdown n=4, all other cells n=5")

    # 10 — close
    slide = new_slide(
        "Three takeaways—and the boundary",
        "Paper conclusion and evidence policy",
        "Close: The contribution is not an experiment catalogue. It is one execution key used to expose collapse, retain discovered support, and test a verified discovery extension.",
    )
    add_title(slide, "Three takeaways—and the boundary", 10, subtitle="The narrow version of the claim is the memorable—and defensible—version.")
    card_x = [0.58, 4.77, 8.96]
    card_fill = [TEAL_LIGHT, CARD, PURPLE_LIGHT]
    card_num = ["01", "02", "03"]
    card_head = ["MEASURE", "RETAIN", "DISCOVER"]
    card_body = [
        "Executable keys reveal support collapse that binary correctness and token entropy miss.",
        "Uniform verified replay raises raw breadth in all 75 / 75 primary pairs; excess breadth is domain-dependent.",
        "Verified-support pressure adds positive mean adjusted breadth at two scales in an exploratory 49-endpoint bundle.",
    ]
    card_color = [TEAL_DARK, INK, PURPLE]
    for x, fill, num, head, body, color in zip(card_x, card_fill, card_num, card_head, card_body, card_color):
        add_card(slide, x, 1.56, 3.78, 3.72, fill=fill)
        add_text(slide, num, x + 0.26, 1.88, 0.62, 0.40, size=20, color=color, bold=True)
        add_text(slide, head, x + 0.93, 1.91, 2.45, 0.30, size=13, color=color, bold=True)
        add_text(slide, body, x + 0.28, 2.60, 3.20, 1.86, size=17, color=INK, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_card(slide, 0.76, 5.61, 11.82, 1.12, fill=INK)
    add_rich_text(slide, [{"text": "CLAIM  ", "bold": True, "color": TEAL_LIGHT}, {"text": "Executed mode identity exposes hidden collapse; verified replay retains discovered modes.", "bold": True, "color": WHITE}], 1.04, 5.84, 11.25, 0.34, size=14.3, align=PP_ALIGN.CENTER)
    add_text(slide, "BOUNDARY  No automatic utility · replay cannot recover unseen modes · synthetic exact-mode tasks · discovery is bundled and exploratory", 1.08, 6.28, 11.18, 0.25, size=9.4, color=CORAL_LIGHT, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "Stop here for the core talk · the remaining slides are backup")

    # 11 — backup: direct alternatives
    slide = new_slide(
        "Only replay is uniform; alternatives are local",
        str(DIRECT_JSON.relative_to(ROOT)),
        "Backup: Use if asked whether ordinary GRPO, UCPO, or generic verified-success replay reproduces the effect. Complete smaller-model adjusted-breadth blocks are heterogeneous.",
    )
    add_title(slide, "Only replay is uniform; alternatives are local", 11, kicker="BACKUP · DIRECT ALTERNATIVES", subtitle="Complete smaller-model correctness-adjusted breadth effects against matched Dr.GRPO.")
    add_picture_contain(slide, plots["comparators"], 0.53, 1.40, 12.28, 5.46)
    add_source(slide, "GRPO, UCPO, and RLEP-Dr are heterogeneous; Falcon Python RLEP-Dr is n=2 and omitted from the n=5 heatmap")

    # 12 — backup: trajectory
    slide = new_slide(
        "Where excess breadth appears, it persists through training",
        str(AUC_JSON.relative_to(ROOT)),
        "Backup: Full-trajectory normalized AUC confirms that Graph, Countdown, and Pantry effects are not terminal-checkpoint accidents at Qwen2.5-0.5B.",
    )
    add_title(slide, "Where excess breadth appears, it persists through training", 12, kicker="BACKUP · TRAJECTORY", subtitle="Qwen2.5-0.5B normalized trajectory AUC; five paired seeds per domain.")
    add_picture_contain(slide, plots["auc"], 0.54, 1.42, 12.26, 5.48)
    add_source(slide, "Graph, Countdown, and Pantry adjusted-breadth AUC intervals exclude zero; Python and MathIR track accuracy rescue")

    # 13 — backup: telemetry
    slide = new_slide(
        "The retention mechanism is small and measurable",
        str(TELEMETRY_JSON.relative_to(ROOT)),
        "Backup: The replay bank is prompt-local and small. Capacity rarely binds. Optimizer timing is descriptive, not an inferential compute claim.",
    )
    add_title(slide, "The retention mechanism is small and measurable", 13, kicker="BACKUP · MECHANISM TELEMETRY", subtitle="Prompt-local banks are continuously revisited; they are not a giant offline replay corpus.")
    add_picture_contain(slide, plots["telemetry"], 0.53, 1.39, 12.30, 5.55)
    add_source(slide, "Qwen2.5-0.5B telemetry · 25 replay + 25 exact-zero control logs · timing is descriptive")

    return prs, manifest_slides

def validate_pptx(path: Path, expected_slides: int) -> dict:
    with zipfile.ZipFile(path) as zf:
        bad = zf.testzip()
        if bad:
            raise RuntimeError(f"Corrupt PPTX member: {bad}")
    deck = Presentation(path)
    if len(deck.slides) != expected_slides:
        raise RuntimeError(f"Expected {expected_slides} slides, found {len(deck.slides)}")
    overflows = []
    for sidx, slide in enumerate(deck.slides, start=1):
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > deck.slide_width + 2 or shape.top + shape.height > deck.slide_height + 2:
                overflows.append({"slide": sidx, "shape": shape.name})
    return {"slides": len(deck.slides), "zip_ok": True, "shape_overflows": overflows, "bytes": path.stat().st_size}


def main() -> None:
    FIG.mkdir(parents=True, exist_ok=True)
    setup_matplotlib()
    cross = read_json(CROSS_JSON)
    auc = read_json(AUC_JSON)
    telemetry = read_json(TELEMETRY_JSON)
    direct = read_json(DIRECT_JSON)
    discovery = read_json(DISCOVERY_JSON)
    assets = copy_conceptual_assets()
    plots = {
        "headline": build_headline_pairs(cross),
        "scatter": build_accuracy_breadth_scatter(cross),
        "discovery": build_discovery_effects(discovery),
        "auc": build_auc_effects(auc),
        "telemetry": build_telemetry(telemetry),
        "comparators": build_comparator_heatmap(cross, direct),
    }
    prs, slide_manifest = build_pptx(assets, plots, cross, discovery)
    prs.save(PPTX_OUT)
    validation = validate_pptx(PPTX_OUT, 13)
    evidence_files = [CROSS_JSON, DISCOVERY_JSON, AUC_JSON, TELEMETRY_JSON, DIRECT_JSON, PAPER_COLLAPSE, PAPER_MODEBENCH, PAPER_METHOD]
    manifest = {
        "deck": str(PPTX_OUT.relative_to(ROOT)),
        "storyboard": "presentation/STORYBOARD.md",
        "slides": slide_manifest,
        "validation": validation,
        "evidence_sha256": {str(p.relative_to(ROOT)): sha256(p) for p in evidence_files},
        "derived_figures": {k: str(v.relative_to(ROOT)) for k, v in plots.items()},
        "claim_boundary": "Retention is the complete 75-pair core; discovery is a 49-endpoint exploratory bundled contrast; adjusted breadth separates added modes from added correctness.",
    }
    MANIFEST_OUT.write_text(json.dumps(manifest, indent=2) + "\n")
    print(json.dumps({"pptx": str(PPTX_OUT), "manifest": str(MANIFEST_OUT), "validation": validation}, indent=2))


if __name__ == "__main__":
    main()
